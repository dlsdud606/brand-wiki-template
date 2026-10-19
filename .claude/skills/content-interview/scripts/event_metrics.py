#!/usr/bin/env python3
"""Read the event source workbook; keep its denominators separate from calendar days."""
import argparse
import hashlib
import json
import math
import posixpath
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

NS = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'p': 'http://schemas.openxmlformats.org/package/2006/relationships'}
HEADERS = ['카테고리', '이벤트명', '이벤트조회수', '이벤트시작', '이벤트종료', '게시기간', '일평균', '링크']

def number(value):
    if isinstance(value, bool) or value is None:
        return None
    try:
        n = float(value)
    except (TypeError, ValueError):
        return None
    return n if math.isfinite(n) else None

def day(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip().replace('.', '-'))
        except ValueError:
            pass
    return None

def metrics(values):
    views, duration = number(values[2]), number(values[5])
    cached = number(values[6])
    start, end = day(values[3]), day(values[4])
    warnings = []
    if views is None or views < 0 or not views.is_integer():
        warnings.append('조회수 누락/음수/비정수: 일평균 계산 보류')
        views = None
    if duration is None or duration <= 0:
        warnings.append('원본 게시기간 누락/0 이하: 원본 분모 일평균 보류')
        duration = None
    if not start or not end or end < start:
        warnings.append('날짜 누락/역전: 달력일수 계산 보류')
        delta, inclusive = None, None
    else:
        delta, inclusive = (end - start).days, (end - start).days + 1
        if duration is not None and duration != delta:
            warnings.append('원본 게시기간과 종료−시작일 불일치: 원본 분모는 보존')
    daily = views / duration if views is not None and duration is not None else None
    calendar_daily = views / inclusive if views is not None and inclusive is not None else None
    if cached is None:
        warnings.append('원본 일평균 캐시 없음: 재계산값 별도 기록')
    elif daily is not None and abs(cached - daily) > 0.01:
        warnings.append('원본 일평균 캐시와 조회수/게시기간 불일치')
    return {
        'date_difference_days': delta, 'inclusive_calendar_days': inclusive,
        'daily_source_denominator': daily, 'daily_inclusive_calendar': calendar_daily,
        'warnings': warnings,
    }

def workbook_rows_stdlib(source):
    """Read cached XLSX values when openpyxl is unavailable; never evaluate formulas."""
    with ZipFile(source) as archive:
        workbook = ElementTree.fromstring(archive.read('xl/workbook.xml'))
        relationships = ElementTree.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {rel.attrib['Id']: rel.attrib['Target'] for rel in relationships.findall('p:Relationship', NS)}
        strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            shared = ElementTree.fromstring(archive.read('xl/sharedStrings.xml'))
            strings = [''.join(node.itertext()) for node in shared.findall('x:si', NS)]
        styles = ElementTree.fromstring(archive.read('xl/styles.xml'))
        date_formats = {int(node.attrib['numFmtId']) for node in styles.findall('x:numFmts/x:numFmt', NS)
                        if any(token in node.attrib.get('formatCode', '').lower() for token in ('yy', 'dd', 'mm'))}
        cell_styles = [int(node.attrib.get('numFmtId', '0')) for node in styles.findall('x:cellXfs/x:xf', NS)]
        properties = workbook.find('x:workbookPr', NS)
        epoch = datetime(1904, 1, 1) if properties is not None and properties.attrib.get('date1904') == '1' else datetime(1899, 12, 30)
        for sheet in workbook.findall('x:sheets/x:sheet', NS):
            target = targets[sheet.attrib[f'{{{NS["r"]}}}id']]
            path = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl', target))
            root = ElementTree.fromstring(archive.read(path))
            for row in root.findall('x:sheetData/x:row', NS):
                values = [None] * 8
                for cell in row.findall('x:c', NS):
                    match = re.match(r'^([A-H])\d+$', cell.attrib.get('r', ''))
                    if not match:
                        continue
                    value = cell.find('x:v', NS)
                    cell_type = cell.attrib.get('t')
                    if cell_type == 'inlineStr':
                        inline = cell.find('x:is', NS)
                        parsed = ''.join(inline.itertext()) if inline is not None else None
                    elif value is None or value.text is None:
                        parsed = None
                    elif cell_type == 's':
                        parsed = strings[int(value.text)]
                    elif cell_type in ('str', 'e'):
                        parsed = value.text
                    else:
                        parsed = float(value.text)
                        style = int(cell.attrib.get('s', '0'))
                        format_id = cell_styles[style]
                        if format_id in set(range(14, 23)) | date_formats:
                            parsed = epoch + timedelta(days=parsed)
                    values[ord(match.group(1)) - ord('A')] = parsed
                yield sheet.attrib['name'], int(row.attrib['r']), values

def source_rows(source):
    try:
        from openpyxl import load_workbook
    except ImportError:
        yield from workbook_rows_stdlib(source)
        return
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        for sheet in book:
            for rownum, vals in enumerate(sheet.iter_rows(values_only=True), 1):
                yield sheet.title, rownum, list(vals[:8]) + [None] * max(0, 8 - len(vals))
    finally:
        book.close()

def extract(source):
    rows = []
    seen_headers = set()
    for sheet, rownum, vals in source_rows(source):
        if sheet not in seen_headers:
            if not any(v is not None for v in vals):
                continue
            if list(vals[:8]) != HEADERS:
                raise ValueError(f'{sheet}: 헤더 변경. 실제 필드를 확인한 뒤 추출기를 수정하세요.')
            seen_headers.add(sheet)
            continue
        if not any(v is not None for v in vals):
            continue
        raw = dict(zip(HEADERS, [v.isoformat() if isinstance(v, (date, datetime)) else v for v in vals]))
        rows.append({'source_sheet': sheet, 'source_row': rownum, 'source_values': raw, **metrics(vals)})
    if not rows:
        raise ValueError('비교할 원본 행이 없습니다.')
    return {
        'source_file': str(source), 'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
        'metric': '이벤트 페이지 조회수. 순방문자/구매/매출 아님',
        'denominator_policy': '원본 게시기간과 양끝 날짜 포함 달력일수를 각각 계산. 원본 미수정.',
        'row_count': len(rows),
        'total_valid_views': sum(
            n for r in rows
            if (n := number(r['source_values']['이벤트조회수'])) is not None
            and n >= 0 and n.is_integer()
        ),
        'rows': rows,
    }

def main():
    wiki = Path(__file__).resolve().parents[4]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=wiki / 'raw/reports/이벤트_조회수.xlsx')
    parser.add_argument('--out', type=Path, help='JSON output inside this wiki deliverables; otherwise stdout')
    args = parser.parse_args()
    if args.out:
        destination = args.out.resolve()
        allowed = (wiki / 'deliverables').resolve()
        if not destination.is_relative_to(allowed):
            parser.error('--out은 위키 deliverables/ 내부여야 합니다.')
        if destination == args.source.resolve():
            parser.error('원본에 쓰지 않습니다.')
    data = extract(args.source.resolve())
    result = json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False)
    if args.out:
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(result + '\n', encoding='utf-8')
        print(f'{data["row_count"]}행 추출: {destination}')
    else:
        print(result)

if __name__ == '__main__':
    main()
