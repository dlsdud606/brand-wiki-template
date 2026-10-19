#!/usr/bin/env python3
"""레퍼런스 릴스 엑셀(raw/reference-reels/*.xlsx)을 읽기 전용으로 요약한다.

사용:
  python3 reference_grounding.py <xlsx> [--type A,B,F] [--top 5] [--kr-only] [--out 요약.md]

출력(마크다운): 유형별 편수·국내/해외·최고 조회수·대표 사례 표, 유형별 상위 N편 표(브랜드·국가·게시일·조회수·영상 내용·링크).
썸네일·카피는 옮기지 않는다. 조회수는 수집 시점 표시값이며 노출 경로를 알 수 없다는 고지를 항상 붙인다.
"""
import argparse, sys
from collections import defaultdict

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl이 필요합니다: python3 -m pip install openpyxl")

def load(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["레퍼런스_70"] if "레퍼런스_70" in wb.sheetnames else wb[wb.sheetnames[-1]]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    idx = {h: i for i, h in enumerate(header)}
    need = ["유형", "No", "브랜드", "국가", "게시 계정", "게시일", "조회수", "영상 내용", "릴스 링크"]
    missing = [n for n in need if n not in idx]
    if missing:
        sys.exit(f"열이 없습니다: {missing} / 실제 열: {header}")
    out = []
    for r in rows[1:]:
        if r is None or r[idx["유형"]] is None:
            continue
        rec = {n: r[idx[n]] for n in need}
        try:
            rec["조회수"] = int(rec["조회수"] or 0)
        except (TypeError, ValueError):
            rec["조회수"] = 0
        rec["type_code"] = str(rec["유형"]).strip()[0]
        rec["collab"] = "(협업)" in str(rec["게시 계정"] or "")
        d = rec["게시일"]
        rec["게시일"] = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else (str(d) if d else "")
        out.append(rec)
    return out

def fmt_views(v):
    return f"{v/10000:.1f}만" if v >= 10000 else f"{v:,}"

def summary(recs):
    by = defaultdict(list)
    for r in recs:
        by[r["type_code"]].append(r)
    lines = ["| 유형 | 편수 | 국내/해외 | 최고 조회수 | 대표 사례 | 협업 편수 |", "|---|---|---|---|---|---|"]
    for code in sorted(by):
        rs = sorted(by[code], key=lambda x: -x["조회수"])
        kr = sum(1 for r in rs if str(r["국가"]).upper() == "KR")
        top = rs[0]
        lines.append(
            f"| {rs[0]['유형']} | {len(rs)} | {kr} / {len(rs)-kr} | {fmt_views(top['조회수'])} | {top['브랜드']} ({top['게시일']}) | {sum(r['collab'] for r in rs)} |"
        )
    return "\n".join(lines)

def detail(recs, types, top, kr_only):
    by = defaultdict(list)
    for r in recs:
        if types and r["type_code"] not in types:
            continue
        if kr_only and str(r["국가"]).upper() != "KR":
            continue
        by[r["type_code"]].append(r)
    blocks = []
    for code in sorted(by):
        rs = sorted(by[code], key=lambda x: -x["조회수"])[:top]
        blocks.append(f"\n### {rs[0]['유형']}\n")
        blocks.append("| No | 브랜드 | 국가 | 게시 계정 | 게시일 | 조회수 | 영상 내용 (형식 참고용) | 링크 |")
        blocks.append("|---|---|---|---|---|---|---|---|")
        for r in rs:
            blocks.append(
                f"| {code}-{r['No']} | {r['브랜드']} | {r['국가']} | {r['게시 계정']} | {r['게시일']} | {r['조회수']:,} | {r['영상 내용']} | {r['릴스 링크']} |"
            )
    return "\n".join(blocks)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--type", default="", help="예: A,B,F")
    ap.add_argument("--top", type=int, default=5)
    ap.add_argument("--kr-only", action="store_true")
    ap.add_argument("--out")
    a = ap.parse_args()
    recs = load(a.xlsx)
    types = [t.strip().upper() for t in a.type.split(",") if t.strip()]
    md = [
        f"# 레퍼런스 릴스 요약 (읽기 전용 추출)",
        f"- 원본: `{a.xlsx}` · 총 {len(recs)}편",
        "- 조회수는 수집 시점 릴스 탭 표시값. 노출 경로·광고 집행 여부는 알 수 없음. 유형 분류는 썸네일·캡션 기준",
        "- 타사 공개 게시물. **유형과 영상 표현 방식만 참고**하고 화면·카피·소재는 옮기지 않음",
        "",
        "## 유형별 분포",
        summary(recs),
        "",
        f"## 유형별 상위 {a.top}편" + (f" (유형 {','.join(types)})" if types else "") + (" (국내만)" if a.kr_only else ""),
        detail(recs, types, a.top, a.kr_only),
    ]
    text = "\n".join(md) + "\n"
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"저장: {a.out}")
    else:
        print(text)

if __name__ == "__main__":
    main()
